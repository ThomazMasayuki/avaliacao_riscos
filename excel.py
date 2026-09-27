"""Exportação da matriz de riscos para Excel (.xlsx) com fórmulas vivas.

O Excel é gerado para REVISÃO DO ANALISTA:
  * as notas lançadas na avaliação (P, I, desenho, operação) ficam em azul e
    podem ser ajustadas pelo analista — todos os cálculos se refazem sozinhos;
  * todos os resultados (inerente, efetividade, residual, níveis) são fórmulas
    que leem a aba "Parâmetros" (espelho de config/metodologia.yaml);
  * as colunas de parecer (fundo amarelo) são preenchidas pelo analista.
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .config import Metodologia
from .servico import RiscoAvaliado

FONTE = "Arial"
AZUL_ESCURO = "1F3A5F"
CINZA_CLARO = "F2F4F7"
AMARELO = "FFF4C2"

F_TITULO = Font(name=FONTE, size=16, bold=True, color=AZUL_ESCURO)
F_SUBTIT = Font(name=FONTE, size=11, bold=True, color=AZUL_ESCURO)
F_CAB = Font(name=FONTE, size=10, bold=True, color="FFFFFF")
F_NORMAL = Font(name=FONTE, size=10)
F_INPUT = Font(name=FONTE, size=10, color="0000FF")
F_NOTA = Font(name=FONTE, size=9, italic=True, color="555555")

P_CAB = PatternFill("solid", fgColor=AZUL_ESCURO)
P_CAB_ANALISTA = PatternFill("solid", fgColor="B7791F")
P_AMARELO = PatternFill("solid", fgColor=AMARELO)
P_CINZA = PatternFill("solid", fgColor=CINZA_CLARO)

FINO = Side(style="thin", color="C9CED6")
BORDA = Border(left=FINO, right=FINO, top=FINO, bottom=FINO)
QUEBRA = Alignment(wrap_text=True, vertical="top")
CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)

PARECERES = ["Concordo", "Concordo com ajuste", "Discordo", "Pendente de evidência"]


def _hex(cor: str) -> str:
    return cor.lstrip("#").upper()


def _cabecalho(ws, linha: int, colunas: list[tuple[str, int]], analista_de: int | None = None):
    for i, (nome, largura) in enumerate(colunas, start=1):
        c = ws.cell(row=linha, column=i, value=nome)
        c.font = F_CAB
        c.fill = P_CAB_ANALISTA if analista_de and i >= analista_de else P_CAB
        c.alignment = CENTRO
        c.border = BORDA
        ws.column_dimensions[get_column_letter(i)].width = largura
    ws.row_dimensions[linha].height = 32


def _celula(ws, linha, coluna, valor, fonte=F_NORMAL, fmt=None, alinh=QUEBRA, preench=None):
    c = ws.cell(row=linha, column=coluna, value=valor)
    c.font, c.alignment, c.border = fonte, alinh, BORDA
    if fmt:
        c.number_format = fmt
    if preench:
        c.fill = preench
    return c


def _formatar_niveis(ws, faixa_celulas: str, m: Metodologia):
    """Formatação condicional por texto do nível (Baixo, Médio, Alto, Crítico)."""
    for f in m.faixas_risco:
        cor = _hex(f.cor)
        fonte_cor = "FFFFFF" if f.rotulo in ("Crítico", "Alto") else "000000"
        ws.conditional_formatting.add(faixa_celulas, CellIsRule(
            operator="equal", formula=[f'"{f.rotulo}"'],
            fill=PatternFill("solid", fgColor=cor, bgColor=cor),
            font=Font(name=FONTE, bold=True, color=fonte_cor)))


# =============================================================================
# Aba Parâmetros
# =============================================================================
class _Params:
    """Guarda os endereços absolutos das tabelas de parâmetros."""


def _aba_parametros(wb, m: Metodologia) -> _Params:
    ws = wb.create_sheet("Parâmetros")
    ws["A1"] = "Parâmetros da metodologia"
    ws["A1"].font = F_TITULO
    ws["A2"] = (f"Versão {m.versao} — espelho de config/metodologia.yaml. Alterar aqui "
                "recalcula a matriz inteira (útil para simulações do analista).")
    ws["A2"].font = F_NOTA
    ref = _Params()
    ab = "'Parâmetros'!"

    def tabela(linha, titulo, colunas, linhas, larguras=None):
        ws.cell(row=linha, column=1, value=titulo).font = F_SUBTIT
        _cabecalho_local(linha + 1, colunas)
        for r, valores in enumerate(linhas, start=linha + 2):
            for c, v in enumerate(valores, start=1):
                fonte = F_INPUT if isinstance(v, (int, float)) and c > 1 else F_NORMAL
                _celula(ws, r, c, v, fonte=fonte,
                        fmt="0.00" if isinstance(v, float) else None)
        return linha + 2, linha + 1 + len(linhas)

    def _cabecalho_local(linha, colunas):
        for i, nome in enumerate(colunas, start=1):
            c = ws.cell(row=linha, column=i, value=nome)
            c.font, c.fill, c.alignment, c.border = F_CAB, P_CAB, CENTRO, BORDA

    for col, w in zip("ABCDE", (14, 26, 80, 12, 14)):
        ws.column_dimensions[col].width = w

    linha = 4
    ini, fim = tabela(linha, "Probabilidade", ["Nível", "Rótulo", "Descrição"],
                      [(n.valor, n.rotulo, n.descricao) for n in m.probabilidade.values()])
    ref.prob_nivel, ref.prob_rotulo = f"{ab}$A${ini}:$A${fim}", f"{ab}$B${ini}:$B${fim}"
    linha = fim + 2

    ini, fim = tabela(linha, "Impacto", ["Nível", "Rótulo", "Descrição"],
                      [(n.valor, n.rotulo, n.descricao) for n in m.impacto.values()])
    ref.imp_nivel, ref.imp_rotulo = f"{ab}$A${ini}:$A${fim}", f"{ab}$B${ini}:$B${fim}"
    linha = fim + 2

    ini, fim = tabela(linha, "Controle — Desenho", ["Nível", "Rótulo", "Descrição", "Fator"],
                      [(n.valor, n.rotulo, n.descricao, float(n.fator)) for n in m.desenho.values()])
    ref.des_nivel, ref.des_fator = f"{ab}$A${ini}:$A${fim}", f"{ab}$D${ini}:$D${fim}"
    linha = fim + 2

    ini, fim = tabela(linha, "Controle — Operação", ["Nível", "Rótulo", "Descrição", "Fator"],
                      [(n.valor, n.rotulo, n.descricao, float(n.fator)) for n in m.operacao.values()])
    ref.op_nivel, ref.op_fator = f"{ab}$A${ini}:$A${fim}", f"{ab}$D${ini}:$D${fim}"
    linha = fim + 2

    tipos = list(m.tipos_controle.values())
    ini, fim = tabela(linha, "Tipos de controle", ["Código", "Rótulo", "Natureza", "Peso"],
                      [(t.codigo, t.rotulo, t.natureza, float(t.peso)) for t in tipos])
    ref.tipo_cod = f"{ab}$A${ini}:$A${fim}"
    ref.tipo_rot = f"{ab}$B${ini}:$B${fim}"
    ref.tipo_nat = f"{ab}$C${ini}:$C${fim}"
    ref.tipo_peso = f"{ab}$D${ini}:$D${fim}"
    ref.tipo_lista_abs = f"{ab}$A${ini}:$A${fim}"
    linha = fim + 2

    # Faixas: ordenadas por mínimo (MATCH aproximado)
    faixas = sorted(m.faixas_risco, key=lambda f: f.minimo)
    ini, fim = tabela(linha, "Faixas de nível de risco (pontuação 1 a 25)",
                      ["Mínimo", "Rótulo", "Máximo"],
                      [(int(f.minimo), f.rotulo, int(f.maximo)) for f in faixas])
    ref.faixa_min, ref.faixa_rot = f"{ab}$A${ini}:$A${fim}", f"{ab}$B${ini}:$B${fim}"
    ref.faixa_lista_abs = f"{ab}$B${ini}:$B${fim}"
    for r, f in zip(range(ini, fim + 1), faixas):
        ws.cell(row=r, column=2).fill = PatternFill("solid", fgColor=_hex(f.cor))
    linha = fim + 2

    cls = sorted(m.classificacao_controle, key=lambda f: f.minimo)
    ini, fim = tabela(linha, "Classificação da efetividade do controle (mínimo)",
                      ["Mínimo", "Rótulo"], [(float(f.minimo), f.rotulo) for f in cls])
    ref.cls_min, ref.cls_rot = f"{ab}$A${ini}:$A${fim}", f"{ab}$B${ini}:$B${fim}"
    linha = fim + 2

    ws.cell(row=linha, column=1, value="Cálculo").font = F_SUBTIT
    _celula(ws, linha + 1, 1, "Método do residual")
    _celula(ws, linha + 1, 2, m.metodo_residual, fonte=F_INPUT)
    _celula(ws, linha + 1, 3, "proporcional: a efetividade combinada reduz P e I | "
                              "natureza: preventivos reduzem P, detectivos reduzem I")
    _celula(ws, linha + 2, 1, "Efetividade máxima")
    _celula(ws, linha + 2, 2, m.efetividade_maxima, fonte=F_INPUT, fmt="0%")
    _celula(ws, linha + 2, 3, "Teto da efetividade combinada: nenhum conjunto de controles zera o risco.")
    ref.metodo = f"{ab}$B${linha + 1}"
    ref.teto = f"{ab}$B${linha + 2}"
    dv = DataValidation(type="list", formula1='"proporcional,natureza"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(f"B{linha + 1}")
    linha += 4

    ini, fim = tabela(linha, "Pareceres do analista", ["Parecer"], [(p,) for p in PARECERES])
    ref.parecer_lista_abs = f"{ab}$A${ini}:$A${fim}"
    return ref


# =============================================================================
# Aba Controles
# =============================================================================
COLS_CONTROLE = [
    ("Cód. controle", 12), ("Cód. risco", 10), ("Descrição do controle", 42),
    ("Tipo", 8), ("Tipo (descrição)", 22), ("Natureza", 12), ("Frequência", 16),
    ("Responsável", 18), ("Evidência", 26),
    ("Desenho (1-5)", 10), ("Operação (1-5)", 10), ("Amostra", 9), ("Falhas", 8),
    ("% falha", 9), ("Fator desenho", 10), ("Fator operação", 10), ("Peso tipo", 9),
    ("Efetividade", 11), ("Classificação", 18),
    ("aux ln total", 10), ("aux ln prev.", 10), ("aux ln detec.", 10),
    ("Observação do avaliador", 30),
    ("Parecer do analista", 20), ("Comentário do analista", 36),
]
C = {nome: get_column_letter(i) for i, (nome, _) in enumerate(COLS_CONTROLE, start=1)}


def _aba_controles(wb, itens: list[RiscoAvaliado], m: Metodologia, p: _Params) -> tuple[int, int]:
    ws = wb.create_sheet("Controles")
    ws["A1"] = "Avaliação dos controles"
    ws["A1"].font = F_TITULO
    ws["A2"] = ("Azul = nota lançada na avaliação (o analista pode ajustar). "
                "Preto = fórmula. Colunas laranja = preenchimento do analista.")
    ws["A2"].font = F_NOTA
    cab = 4
    idx_analista = [n for n, _ in COLS_CONTROLE].index("Parecer do analista") + 1
    _cabecalho(ws, cab, COLS_CONTROLE, analista_de=idx_analista)

    linha = cab + 1
    for it in itens:
        for c in it.controles:
            L = linha
            vals = {
                "Cód. controle": (c["codigo"], F_NORMAL, None),
                "Cód. risco": (it.risco["codigo"], F_NORMAL, None),
                "Descrição do controle": (c["descricao"], F_NORMAL, None),
                "Tipo": (c["tipo"], F_INPUT, None),
                "Tipo (descrição)": (f'=IFERROR(INDEX({p.tipo_rot},MATCH({C["Tipo"]}{L},{p.tipo_cod},0)),"")', F_NORMAL, None),
                "Natureza": (f'=IFERROR(INDEX({p.tipo_nat},MATCH({C["Tipo"]}{L},{p.tipo_cod},0)),"")', F_NORMAL, None),
                "Frequência": (c.get("frequencia") or "", F_NORMAL, None),
                "Responsável": (c.get("responsavel") or "", F_NORMAL, None),
                "Evidência": (c.get("evidencia") or "", F_NORMAL, None),
                "Desenho (1-5)": (c["desenho"], F_INPUT, "0"),
                "Operação (1-5)": (c["operacao"], F_INPUT, "0"),
                "Amostra": (c.get("amostra"), F_INPUT, "0"),
                "Falhas": (c.get("falhas"), F_INPUT, "0"),
                "% falha": (f'=IF(N({C["Amostra"]}{L})>0,{C["Falhas"]}{L}/{C["Amostra"]}{L},"")', F_NORMAL, "0.0%"),
                "Fator desenho": (f'=INDEX({p.des_fator},MATCH({C["Desenho (1-5)"]}{L},{p.des_nivel},0))', F_NORMAL, "0.00"),
                "Fator operação": (f'=INDEX({p.op_fator},MATCH({C["Operação (1-5)"]}{L},{p.op_nivel},0))', F_NORMAL, "0.00"),
                "Peso tipo": (f'=INDEX({p.tipo_peso},MATCH({C["Tipo"]}{L},{p.tipo_cod},0))', F_NORMAL, "0.00"),
                "Efetividade": (f'={C["Fator desenho"]}{L}*{C["Fator operação"]}{L}*{C["Peso tipo"]}{L}', F_NORMAL, "0.0%"),
                "Classificação": (f'=INDEX({p.cls_rot},MATCH({C["Efetividade"]}{L}+0.000000001,{p.cls_min},1))', F_NORMAL, None),
                "aux ln total": (f'=LN(MAX(1-{C["Efetividade"]}{L},0.000000001))', F_NORMAL, "0.0000"),
                "aux ln prev.": (f'=IF({C["Natureza"]}{L}="preventivo",{C["aux ln total"]}{L},0)', F_NORMAL, "0.0000"),
                "aux ln detec.": (f'=IF({C["Natureza"]}{L}="detectivo",{C["aux ln total"]}{L},0)', F_NORMAL, "0.0000"),
                "Observação do avaliador": (c.get("observacao") or "", F_NORMAL, None),
                "Parecer do analista": (None, F_NORMAL, None),
                "Comentário do analista": (None, F_NORMAL, None),
            }
            for col_i, (nome, _) in enumerate(COLS_CONTROLE, start=1):
                v, fonte, fmt = vals[nome]
                preench = P_AMARELO if col_i >= idx_analista else None
                al = CENTRO if fmt or nome in ("Tipo", "Classificação", "Natureza") else QUEBRA
                _celula(ws, L, col_i, v, fonte=fonte, fmt=fmt, alinh=al, preench=preench)
            linha += 1

    ultima = max(linha - 1, cab + 1)
    ws.freeze_panes = ws.cell(row=cab + 1, column=4)
    ws.auto_filter.ref = f"A{cab}:{get_column_letter(len(COLS_CONTROLE))}{ultima}"

    # Colunas auxiliares agrupadas (ocultas, mas expansíveis)
    ws.column_dimensions.group(C["aux ln total"], C["aux ln detec."], hidden=True)

    # Validações e formatação
    dv_escala = DataValidation(type="whole", operator="between", formula1="1", formula2="5",
                               showErrorMessage=True, errorTitle="Nota inválida",
                               error="Informe um número inteiro de 1 a 5.")
    dv_tipo = DataValidation(type="list", formula1=p.tipo_lista_abs.replace("'Parâmetros'!", "='Parâmetros'!"))
    dv_parecer = DataValidation(type="list", formula1="=" + p.parecer_lista_abs)
    for dv in (dv_escala, dv_tipo, dv_parecer):
        ws.add_data_validation(dv)
    dv_escala.add(f"{C['Desenho (1-5)']}{cab + 1}:{C['Operação (1-5)']}{ultima + 200}")
    dv_tipo.add(f"{C['Tipo']}{cab + 1}:{C['Tipo']}{ultima + 200}")
    dv_parecer.add(f"{C['Parecer do analista']}{cab + 1}:{C['Parecer do analista']}{ultima + 200}")

    faixa_cls = f"{C['Classificação']}{cab + 1}:{C['Classificação']}{ultima}"
    for f in m.classificacao_controle:
        cor = _hex(f.cor)
        ws.conditional_formatting.add(faixa_cls, CellIsRule(
            operator="equal", formula=[f'"{f.rotulo}"'],
            fill=PatternFill("solid", fgColor=cor, bgColor=cor),
            font=Font(name=FONTE, bold=True, color="FFFFFF")))
    ws.cell(row=cab, column=[n for n, _ in COLS_CONTROLE].index("Efetividade") + 1).comment = Comment(
        "Efetividade = fator de desenho × fator de operação × peso do tipo de controle.", "Metodologia")
    return cab + 1, ultima


# =============================================================================
# Aba Matriz de Riscos
# =============================================================================
COLS_RISCO = [
    ("Cód. risco", 10), ("Processo", 22), ("Categoria", 18), ("Descrição do risco", 40),
    ("Causa", 28), ("Consequência", 28), ("Responsável", 18),
    ("Probabilidade (1-5)", 12), ("Impacto (1-5)", 11), ("Risco inerente", 10), ("Nível inerente", 12),
    ("Qtd. controles", 10), ("Efetividade combinada", 12), ("Efet. preventiva", 11), ("Efet. detectiva", 11),
    ("Prob. residual", 10), ("Impacto residual", 10), ("Risco residual", 10), ("Nível residual", 12),
    ("Redução do risco", 10), ("Posição P residual", 10), ("Posição I residual", 10),
    ("Justificativa do avaliador", 34),
    ("Parecer do analista", 20), ("Nível proposto pelo analista", 16),
    ("Comentário do analista", 36), ("Analista", 16), ("Data da revisão", 12),
]
R = {nome: get_column_letter(i) for i, (nome, _) in enumerate(COLS_RISCO, start=1)}


def _aba_matriz(wb, itens, m, p: _Params, ctrl_ini: int, ctrl_fim: int) -> tuple[int, int]:
    ws = wb.create_sheet("Matriz de Riscos")
    ws["A1"] = "Matriz de riscos"
    ws["A1"].font = F_TITULO
    ws["A2"] = ("Azul = nota da avaliação (ajustável). Preto = fórmula. "
                "Colunas laranja = parecer do analista.")
    ws["A2"].font = F_NOTA
    cab = 4
    idx_analista = [n for n, _ in COLS_RISCO].index("Parecer do analista") + 1
    _cabecalho(ws, cab, COLS_RISCO, analista_de=idx_analista)

    # Faixas da aba Controles (com folga para linhas adicionadas pelo analista)
    fim_c = ctrl_fim + 200
    rng = lambda col: f"Controles!${C[col]}${ctrl_ini}:${C[col]}${fim_c}"

    linha = cab + 1
    for it in itens:
        r, L = it.risco, linha
        f = {
            "Cód. risco": (r["codigo"], F_NORMAL, None),
            "Processo": (r["processo"], F_NORMAL, None),
            "Categoria": (r.get("categoria") or "", F_NORMAL, None),
            "Descrição do risco": (r["descricao"], F_NORMAL, None),
            "Causa": (r.get("causa") or "", F_NORMAL, None),
            "Consequência": (r.get("consequencia") or "", F_NORMAL, None),
            "Responsável": (r.get("responsavel") or "", F_NORMAL, None),
            "Probabilidade (1-5)": (r["probabilidade"], F_INPUT, "0"),
            "Impacto (1-5)": (r["impacto"], F_INPUT, "0"),
            "Risco inerente": (f'={R["Probabilidade (1-5)"]}{L}*{R["Impacto (1-5)"]}{L}', F_NORMAL, "0"),
            "Nível inerente": (f'=INDEX({p.faixa_rot},MATCH(ROUND({R["Risco inerente"]}{L},0),{p.faixa_min},1))', F_NORMAL, None),
            "Qtd. controles": (f'=COUNTIF({rng("Cód. risco")},{R["Cód. risco"]}{L})', F_NORMAL, "0"),
            "Efetividade combinada": (f'=MIN({p.teto},1-EXP(SUMIFS({rng("aux ln total")},{rng("Cód. risco")},{R["Cód. risco"]}{L})))', F_NORMAL, "0.0%"),
            "Efet. preventiva": (f'=MIN({p.teto},1-EXP(SUMIFS({rng("aux ln prev.")},{rng("Cód. risco")},{R["Cód. risco"]}{L})))', F_NORMAL, "0.0%"),
            "Efet. detectiva": (f'=MIN({p.teto},1-EXP(SUMIFS({rng("aux ln detec.")},{rng("Cód. risco")},{R["Cód. risco"]}{L})))', F_NORMAL, "0.0%"),
            "Prob. residual": (f'=1+({R["Probabilidade (1-5)"]}{L}-1)*(1-IF({p.metodo}="natureza",{R["Efet. preventiva"]}{L},{R["Efetividade combinada"]}{L}))', F_NORMAL, "0.00"),
            "Impacto residual": (f'=1+({R["Impacto (1-5)"]}{L}-1)*(1-IF({p.metodo}="natureza",{R["Efet. detectiva"]}{L},{R["Efetividade combinada"]}{L}))', F_NORMAL, "0.00"),
            "Risco residual": (f'={R["Prob. residual"]}{L}*{R["Impacto residual"]}{L}', F_NORMAL, "0.00"),
            "Nível residual": (f'=INDEX({p.faixa_rot},MATCH(ROUND({R["Risco residual"]}{L},0),{p.faixa_min},1))', F_NORMAL, None),
            "Redução do risco": (f'=1-{R["Risco residual"]}{L}/{R["Risco inerente"]}{L}', F_NORMAL, "0%"),
            "Posição P residual": (f'=MAX(1,ROUND({R["Prob. residual"]}{L},0))', F_NORMAL, "0"),
            "Posição I residual": (f'=MAX(1,ROUND({R["Impacto residual"]}{L},0))', F_NORMAL, "0"),
            "Justificativa do avaliador": (r.get("justificativa") or "", F_NORMAL, None),
            "Parecer do analista": (None, F_NORMAL, None),
            "Nível proposto pelo analista": (None, F_NORMAL, None),
            "Comentário do analista": (None, F_NORMAL, None),
            "Analista": (None, F_NORMAL, None),
            "Data da revisão": (None, F_NORMAL, "dd/mm/yyyy"),
        }
        for col_i, (nome, _) in enumerate(COLS_RISCO, start=1):
            v, fonte, fmt = f[nome]
            preench = P_AMARELO if col_i >= idx_analista else None
            al = CENTRO if fmt or nome.startswith(("Nível", "Cód")) else QUEBRA
            _celula(ws, L, col_i, v, fonte=fonte, fmt=fmt, alinh=al, preench=preench)
        linha += 1

    ultima = max(linha - 1, cab + 1)
    ws.freeze_panes = ws.cell(row=cab + 1, column=2)
    ws.auto_filter.ref = f"A{cab}:{get_column_letter(len(COLS_RISCO))}{ultima}"
    ws.column_dimensions.group(R["Posição P residual"], R["Posição I residual"], hidden=True)

    for col in ("Nível inerente", "Nível residual", "Nível proposto pelo analista"):
        _formatar_niveis(ws, f"{R[col]}{cab + 1}:{R[col]}{ultima}", m)
    # destaca riscos sem controle
    ws.conditional_formatting.add(
        f"{R['Qtd. controles']}{cab + 1}:{R['Qtd. controles']}{ultima}",
        CellIsRule(operator="equal", formula=["0"],
                   fill=PatternFill("solid", fgColor="FDE2E1", bgColor="FDE2E1"),
                   font=Font(name=FONTE, bold=True, color="C62828")))

    dv_escala = DataValidation(type="whole", operator="between", formula1="1", formula2="5",
                               showErrorMessage=True, errorTitle="Nota inválida",
                               error="Informe um número inteiro de 1 a 5.")
    dv_parecer = DataValidation(type="list", formula1="=" + p.parecer_lista_abs)
    dv_nivel = DataValidation(type="list", formula1="=" + p.faixa_lista_abs)
    dv_data = DataValidation(type="date", operator="greaterThan", formula1="36526")
    for dv in (dv_escala, dv_parecer, dv_nivel, dv_data):
        ws.add_data_validation(dv)
    dv_escala.add(f"{R['Probabilidade (1-5)']}{cab + 1}:{R['Impacto (1-5)']}{ultima}")
    dv_parecer.add(f"{R['Parecer do analista']}{cab + 1}:{R['Parecer do analista']}{ultima}")
    dv_nivel.add(f"{R['Nível proposto pelo analista']}{cab + 1}:{R['Nível proposto pelo analista']}{ultima}")
    dv_data.add(f"{R['Data da revisão']}{cab + 1}:{R['Data da revisão']}{ultima}")

    ws.cell(row=cab, column=[n for n, _ in COLS_RISCO].index("Efetividade combinada") + 1).comment = Comment(
        "Controles independentes: 1 − Π(1 − efetividade de cada controle), limitado ao teto.", "Metodologia")
    ws.cell(row=cab, column=[n for n, _ in COLS_RISCO].index("Prob. residual") + 1).comment = Comment(
        "P residual = 1 + (P − 1) × (1 − efetividade). Idem para impacto.", "Metodologia")
    return cab + 1, ultima


# =============================================================================
# Aba Resumo (KPIs + mapas de calor + distribuição)
# =============================================================================
def _mapa_calor(ws, lin0, col0, titulo, col_p, col_i, faixa_riscos, m: Metodologia):
    ws.cell(row=lin0, column=col0, value=titulo).font = F_SUBTIT
    ws.merge_cells(start_row=lin0 + 1, start_column=col0, end_row=lin0 + 5, end_column=col0)
    rot = ws.cell(row=lin0 + 1, column=col0, value="PROBABILIDADE")
    rot.font, rot.alignment = Font(name=FONTE, size=9, bold=True, color=AZUL_ESCURO), Alignment(
        text_rotation=90, horizontal="center", vertical="center")
    for k, prob in enumerate(range(5, 0, -1)):
        lin = lin0 + 1 + k
        c = ws.cell(row=lin, column=col0 + 1, value=f"{prob} {m.probabilidade[prob].rotulo}")
        c.font, c.alignment = Font(name=FONTE, size=9, bold=True), Alignment(horizontal="right", vertical="center")
        ws.row_dimensions[lin].height = 38
        for imp in range(1, 6):
            cor = _hex(m.faixa_risco(prob * imp).cor)
            cel = ws.cell(row=lin, column=col0 + 1 + imp,
                          value=f'=COUNTIFS({faixa_riscos(col_p)},{prob},{faixa_riscos(col_i)},{imp})')
            cel.fill = PatternFill("solid", fgColor=cor)
            cel.font = Font(name=FONTE, size=14, bold=True)
            cel.alignment = CENTRO
            cel.border = Border(left=Side(style="medium", color="FFFFFF"), right=Side(style="medium", color="FFFFFF"),
                                top=Side(style="medium", color="FFFFFF"), bottom=Side(style="medium", color="FFFFFF"))
            cel.number_format = '0;-0;""'
    for imp in range(1, 6):
        c = ws.cell(row=lin0 + 6, column=col0 + 1 + imp, value=f"{imp} {m.impacto[imp].rotulo}")
        c.font, c.alignment = Font(name=FONTE, size=9, bold=True), CENTRO
    ws.row_dimensions[lin0 + 6].height = 28
    c = ws.cell(row=lin0 + 7, column=col0 + 2, value="IMPACTO")
    ws.merge_cells(start_row=lin0 + 7, start_column=col0 + 2, end_row=lin0 + 7, end_column=col0 + 6)
    c.font, c.alignment = Font(name=FONTE, size=9, bold=True, color=AZUL_ESCURO), CENTRO


def _aba_resumo(wb, aval: dict, m: Metodologia, r_ini, r_fim, c_ini, c_fim):
    ws = wb.create_sheet("Resumo", 0)
    ws.sheet_view.showGridLines = False
    ws["B2"] = f"Matriz de Riscos — {aval['area']}"
    ws["B2"].font = F_TITULO
    ws["B3"] = (f"Período: {aval['periodo']}   |   Avaliador: {aval['avaliador']}   |   "
                f"Status: {aval['status']}   |   Metodologia: {m.versao}   |   "
                f"Gerado em {datetime.now():%d/%m/%Y %H:%M}")
    ws["B3"].font = F_NOTA
    for col, w in zip("ABCDEFGHIJKLMNOPQ", (2, 4, 16, 13, 13, 13, 13, 13, 3, 4, 16, 13, 13, 13, 13, 13, 3)):
        ws.column_dimensions[col].width = w

    fr = lambda col: f"'Matriz de Riscos'!${R[col]}${r_ini}:${R[col]}${r_fim}"
    fc = lambda col: f"Controles!${C[col]}${c_ini}:${C[col]}${c_fim}"

    # KPIs
    kpis = [
        ("Riscos avaliados", f'=COUNTA({fr("Cód. risco")})', "0"),
        ("Controles avaliados", f'=COUNTA({fc("Cód. controle")})', "0"),
        ("Riscos sem controle", f'=COUNTIF({fr("Qtd. controles")},0)', "0"),
        ("Efetividade média", f'=IFERROR(AVERAGE({fc("Efetividade")}),0)', "0%"),
        ("Inerente médio", f'=IFERROR(AVERAGE({fr("Risco inerente")}),0)', "0.0"),
        ("Residual médio", f'=IFERROR(AVERAGE({fr("Risco residual")}),0)', "0.0"),
    ]
    for k, (rot, formula, fmt) in enumerate(kpis):
        col = 3 + k * 2 if k < 3 else 11 + (k - 3) * 2
        c1 = ws.cell(row=5, column=col, value=rot)
        c1.font = Font(name=FONTE, size=9, bold=True, color="555555")
        c2 = ws.cell(row=6, column=col, value=formula)
        c2.font = Font(name=FONTE, size=20, bold=True, color=AZUL_ESCURO)
        c2.number_format = fmt
        for r in (5, 6):
            ws.cell(row=r, column=col).fill = P_CINZA
            ws.cell(row=r, column=col + 1).fill = P_CINZA
            ws.cell(row=r, column=col).alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.merge_cells(start_row=5, start_column=col, end_row=5, end_column=col + 1)
        ws.merge_cells(start_row=6, start_column=col, end_row=6, end_column=col + 1)
    ws.row_dimensions[6].height = 34

    # Mapas de calor
    _mapa_calor(ws, 9, 2, "Mapa de calor — risco INERENTE",
                "Probabilidade (1-5)", "Impacto (1-5)", fr, m)
    _mapa_calor(ws, 9, 10, "Mapa de calor — risco RESIDUAL",
                "Posição P residual", "Posição I residual", fr, m)

    # Distribuição por nível
    lin = 19
    ws.cell(row=lin, column=2, value="Distribuição por nível de risco").font = F_SUBTIT
    for i, nome in enumerate(("Nível", "Inerente", "Residual", "Variação"), start=3):
        c = ws.cell(row=lin + 1, column=i, value=nome)
        c.font, c.fill, c.alignment, c.border = F_CAB, P_CAB, CENTRO, BORDA
    faixas = sorted(m.faixas_risco, key=lambda f: -f.minimo)
    for k, f in enumerate(faixas):
        L = lin + 2 + k
        _celula(ws, L, 3, f.rotulo, fonte=Font(name=FONTE, bold=True), alinh=CENTRO,
                preench=PatternFill("solid", fgColor=_hex(f.cor)))
        _celula(ws, L, 4, f'=COUNTIF({fr("Nível inerente")},C{L})', fmt="0", alinh=CENTRO)
        _celula(ws, L, 5, f'=COUNTIF({fr("Nível residual")},C{L})', fmt="0", alinh=CENTRO)
        _celula(ws, L, 6, f"=E{L}-D{L}", fmt='+0;-0;0', alinh=CENTRO)
    fim_tab = lin + 1 + len(faixas)
    L = fim_tab + 1
    _celula(ws, L, 3, "Total", fonte=Font(name=FONTE, bold=True), alinh=CENTRO, preench=P_CINZA)
    _celula(ws, L, 4, f"=SUM(D{lin + 2}:D{fim_tab})", fonte=Font(name=FONTE, bold=True), fmt="0", alinh=CENTRO, preench=P_CINZA)
    _celula(ws, L, 5, f"=SUM(E{lin + 2}:E{fim_tab})", fonte=Font(name=FONTE, bold=True), fmt="0", alinh=CENTRO, preench=P_CINZA)
    _celula(ws, L, 6, f"=E{L}-D{L}", fonte=Font(name=FONTE, bold=True), fmt='+0;-0;0', alinh=CENTRO, preench=P_CINZA)

    graf = BarChart()
    graf.type = "col"
    graf.title = "Riscos por nível: inerente × residual"
    graf.y_axis.title = "Qtd. de riscos"
    graf.y_axis.majorGridlines = None
    graf.add_data(Reference(ws, min_col=4, max_col=5, min_row=lin + 1, max_row=fim_tab), titles_from_data=True)
    graf.set_categories(Reference(ws, min_col=3, min_row=lin + 2, max_row=fim_tab))
    for serie, cor in zip(graf.series, ("9AA5B1", AZUL_ESCURO)):
        serie.graphicalProperties.solidFill = cor
        serie.graphicalProperties.line.solidFill = cor
    graf.height, graf.width = 7.5, 15
    ws.add_chart(graf, f"K{lin}")

    # Parecer geral do analista
    L = lin + 17
    ws.cell(row=L, column=2, value="Parecer geral do analista").font = F_SUBTIT
    for rr in range(L + 1, L + 6):
        for cc in range(3, 17):
            ws.cell(row=rr, column=cc).fill = P_AMARELO
    ws.cell(row=L + 1, column=3).alignment = QUEBRA
    ws.merge_cells(start_row=L + 1, start_column=3, end_row=L + 5, end_column=16)
    return ws


def _aba_instrucoes(wb, aval, m):
    ws = wb.create_sheet("Instruções")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 100
    ws["B2"] = "Como revisar esta matriz"
    ws["B2"].font = F_TITULO
    linhas = [
        ("Origem", f"Avaliação de riscos da área {aval['area']} ({aval['periodo']}), "
                   f"conduzida por {aval['avaliador']}. Exportada da aplicação de avaliação."),
        ("Legenda de cores", "Texto AZUL = nota lançada na avaliação (pode ser ajustada). Texto PRETO = fórmula, "
                             "não editar. Fundo AMARELO / cabeçalho LARANJA = campos do analista."),
        ("1. Resumo", "Visão geral: indicadores, mapas de calor inerente e residual, distribuição por nível "
                      "e campo de parecer geral."),
        ("2. Matriz de Riscos", "Um risco por linha. Revise P e I, os níveis calculados e registre parecer, "
                                "nível proposto (se discordar), comentário, nome e data."),
        ("3. Controles", "Um controle por linha, vinculado ao risco pelo código. Revise tipo, notas de desenho e "
                         "operação e registre o parecer. Ao alterar uma nota, a matriz e o resumo se recalculam."),
        ("4. Parâmetros", "Escalas, fatores, pesos e faixas da metodologia. Alterações aqui servem para simulação; "
                          "a metodologia oficial fica no arquivo config/metodologia.yaml da aplicação."),
        ("Fórmulas", "Inerente = P × I.  Efetividade do controle = fator desenho × fator operação × peso do tipo.  "
                     "Efetividade combinada = 1 − Π(1 − efetividade), limitada ao teto.  "
                     "P e I residuais = 1 + (nível − 1) × (1 − efetividade).  Residual = P res × I res."),
        ("Método do residual", f"'{m.metodo_residual}' — proporcional: a efetividade combinada reduz P e I; "
                               "natureza: controles preventivos reduzem P e detectivos reduzem I."),
        ("Colunas ocultas", "Colunas auxiliares (logaritmos e posições residuais) estão agrupadas; use o botão "
                            "'+' acima das colunas para exibi-las."),
    ]
    for k, (a, b) in enumerate(linhas, start=4):
        ca = ws.cell(row=k, column=2, value=a)
        cb = ws.cell(row=k, column=3, value=b)
        ca.font = Font(name=FONTE, bold=True, color=AZUL_ESCURO)
        cb.font = F_NORMAL
        ca.alignment = cb.alignment = QUEBRA
        ws.row_dimensions[k].height = 32


def exportar(caminho: Path | str, aval: dict, itens: list[RiscoAvaliado], m: Metodologia) -> Path:
    caminho = Path(caminho)
    wb = Workbook()
    wb.remove(wb.active)
    p = _aba_parametros(wb, m)
    c_ini, c_fim = _aba_controles(wb, itens, m, p)
    r_ini, r_fim = _aba_matriz(wb, itens, m, p, c_ini, c_fim)
    _aba_resumo(wb, aval, m, r_ini, r_fim, c_ini, c_fim + 200)
    _aba_instrucoes(wb, aval, m)
    wb._sheets = [wb["Resumo"], wb["Matriz de Riscos"], wb["Controles"], wb["Parâmetros"], wb["Instruções"]]
    wb.active = 0
    for ws in wb.worksheets:
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    wb["Resumo"].sheet_properties.tabColor = AZUL_ESCURO
    wb["Matriz de Riscos"].sheet_properties.tabColor = "E53935"
    wb["Controles"].sheet_properties.tabColor = "FB8C00"
    caminho.parent.mkdir(parents=True, exist_ok=True)
    wb.save(caminho)
    return caminho
